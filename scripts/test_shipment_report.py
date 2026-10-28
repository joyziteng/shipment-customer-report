#!/usr/bin/env python3
"""Regression tests using throwaway XLSX workbooks under /tmp."""
import datetime as dt
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

import shipment_report as report

SCRIPT = Path(__file__).with_name('shipment_report.py')


class ShipmentReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='shipment-report-test-', dir='/tmp')
        self.base = Path(self.tmp.name)
        self.customer = self.base / 'customer.xlsx'
        self.shipment = self.base / 'shipment.xlsx'
        self.build_customer()
        self.build_source()

    def tearDown(self):
        self.tmp.cleanup()

    def build_customer(self):
        wb = Workbook()
        totals = wb.active
        totals.title = 'TOTALS'
        totals['A1'] = 'PO#'
        totals['A2'] = 'SENTINEL'
        totals['B2'] = '=1+1'
        totals.sheet_properties.tabColor = 'FF0000'
        po = wb.create_sheet('PO#_260823TP')
        po.append(['PO#', 'Style #', 'Description', 'Color', 'Ttl Units Orders', 'Date Ready for Pickup', 'Ttl CBM'])
        for key, color in [('260823TP', 'Bear'), ('CA260823TP', 'Cappuccino'),
                           ('UK260823TP', 'Cappuccino'), ('260823TP', 'White'),
                           ('260823TP', 'Off White')]:
            po.append([key, 'CBL009', 'Pillow', color, 100, None, 10])
        for forwarder in report.FORWARDERS:
            for kind in ('Ttl Units', '% To Ttl Units', 'CBM'):
                po.cell(1, po.max_column + 1, f'{kind} Shipped 10/12-10/16 (via {forwarder})')
        for r in range(2, 7):
            for c in (8, 10, 11, 13):
                po.cell(r, c, 99)
            for c in (9, 12):
                po.cell(r, c, 0.99).number_format = numbers.FORMAT_PERCENTAGE_00
        po['E7'] = '=SUM(E2:E6)'
        for c in (8, 11):
            po.cell(7, c, f'=SUM({report.get_column_letter(c)}2:{report.get_column_letter(c)}6)')
        po['N2'] = '=1+2'
        for suffix, state in [(' visible', 'visible'), (' hidden', 'hidden')]:
            ws = wb.create_sheet('PO#_260510TP' + suffix)
            ws.sheet_state = state
            ws.append(['PO#', 'Style #', 'Description', 'Color', 'Ttl Units Orders', 'Date', 'Ttl CBM'])
            ws.append(['260510TP', 'CBL026', 'Pillow', 'Blue', 100, None, 10])
            ws.append([None, None, None, None, '=SUM(E2:E2)', None, 10])
        decoy = wb.create_sheet('Notes')
        decoy.append(['PO#', 'Style #', 'Description', 'Color', 'Ttl Units Orders', 'Date', 'Ttl CBM'])
        decoy.append(['260823TP', 'CBL009', 'Pillow', 'Bear', 1, None, 1])
        wb.save(self.customer)

    def build_source(self):
        wb = Workbook()
        wb.active.title = 'unrelated'
        ws = wb.create_sheet(report.SOURCE_TITLE)
        for c, header in report.SOURCE_HEADERS.items():
            ws.cell(1, c, header)
        def add(key, color, day, q, cbm, forwarder='FLEXPORT', note=None, sty='CBL009'):
            r = ws.max_row + 1
            for c, value in {5: sty, 6: f'CloudFreeze Pillow {color}', 7: key, 11: q,
                             14: day, 19: cbm, 25: note, 32: forwarder}.items():
                if value is not None:
                    ws.cell(r, c, value)
        add('PO260823TP', 'Bear', dt.date(2026, 10, 12), 20, 2, note='实际待核')
        add('CA260823TP', 'Cappuccino', dt.date(2026, 10, 16), 30, 3)
        add('UK260823TP', 'Cappuccino', dt.date(2026, 10, 17), 40, 4)
        add('260823TP', 'Off White', dt.date(2026, 10, 13), 25, 2.5)
        add('260823TP', 'White', dt.date(2026, 10, 13), 6, 0.6)
        add('260510TP', 'Blue', dt.date(2026, 10, 13), 5, 0.5, sty='CBL026')
        add('260823TP', 'Bear', dt.date(2026, 10, 13), 0, 0, '预计FLEXPORT-liu')
        add('260824-1TP', 'Bear', dt.date(2026, 10, 13), 9, 0.9)
        add('260823TP', 'Cappuccino', dt.date(2026, 10, 13), 7, 0.7, sty='CBL999')
        add('260823TP', 'Off White', dt.date(2026, 10, 14), 11, None)
        add('260823TP', 'White', dt.date(2026, 10, 14), None, 0.3)
        wb.save(self.shipment)

    def run_report(self, name, *args):
        output = self.base / (name + '.xlsx')
        command = [sys.executable, str(SCRIPT), '--customer', str(self.customer),
                   '--shipment', str(self.shipment), '--output', str(output), *args]
        result = subprocess.run(command, text=True, capture_output=True)
        return result, output

    def test_existing_interval_duplicate_po_and_reconciliation(self):
        result, output = self.run_report('existing', '--start', '2026-10-12', '--end', '2026-10-16')
        self.assertEqual(result.returncode, 0, result.stderr)
        issues = json.loads(output.with_suffix('.issues.json').read_text())
        wb = load_workbook(output, data_only=True)
        po = wb['PO#_260823TP']
        self.assertEqual((po['K2'].value, po['M2'].value), (99, 99))  # unknown forwarder preserves original
        self.assertEqual((po['K3'].value, po['M3'].value), (30, 3))
        self.assertEqual((po['K4'].value, po['M4'].value), (0, 0))  # UK falls outside inclusive end
        self.assertEqual(po['K5'].value, 99)  # quantity unavailable, preserve original
        self.assertAlmostEqual(po['M5'].value, 0.9)  # independent CBM still updates
        self.assertEqual(po['K6'].value, 36)  # quantity updates despite missing CBM
        self.assertEqual(po['M6'].value, 99)  # incomplete CBM preserves original
        self.assertEqual(wb['PO#_260510TP visible']['K2'].value, 5)
        self.assertEqual(wb['PO#_260510TP hidden']['K2'].value, 5)
        self.assertEqual(wb['PO#_260510TP hidden'].sheet_state, 'hidden')
        self.assertEqual(issues['区间来源行数'], 10)
        kinds = {i['类型'] for i in issues['待确认']}
        self.assertTrue({'货代反馈需核对', '货代无法识别', '重复PO明细页', '来源PO无客户明细页', '款号颜色无法唯一匹配'} <= kinds)
        entry = next(x for x in issues['来源PO货代汇总'] if x['PO'] == '260824-1')
        self.assertEqual(entry['数量'], 9)
        detail = next(x for x in issues['明细页核对'] if x['页'] == 'PO#_260823TP')
        recon = next(x for x in detail['按PO货代核对'] if x['PO'] == '260823' and x['货代'] == 'FLEXPORT')
        self.assertEqual(recon['来源数量'], 69)
        self.assertEqual(recon['已归集数量'], 36)
        self.assertEqual(recon['未归集数量'], 33)
        with zipfile.ZipFile(self.customer) as old, zipfile.ZipFile(output) as new:
            for name in old.namelist():
                if name not in {'xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml', 'xl/worksheets/sheet4.xml'}:
                    self.assertEqual(old.read(name), new.read(name), name)
            self.assertEqual(old.read('xl/worksheets/sheet1.xml'), new.read('xl/worksheets/sheet1.xml'))
        formula_wb = load_workbook(output, data_only=False)
        self.assertEqual(formula_wb['TOTALS']['B2'].value, '=1+1')
        self.assertEqual(formula_wb['PO#_260823TP']['N2'].value, '=1+2')
        self.assertEqual(po['H1'].value, 'Ttl Units Shipped 10/12-10/16 (via LAURA)')
        self.assertAlmostEqual(po['K7'].value, sum(po[f'K{r}'].value for r in range(2, 7)))

    def test_partial_overlap_aborts_and_creates_no_output(self):
        result, output = self.run_report('overlap', '--start', '2026-10-15', '--end', '2026-10-17')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('部分交叠', result.stderr)
        self.assertFalse(output.exists())
        self.assertFalse(output.with_suffix('.issues.json').exists())

    def test_new_interval_days_and_unknown_date(self):
        result, output = self.run_report('new', '--days', '1', '--today', '2026-10-18')
        self.assertEqual(result.returncode, 0, result.stderr)
        wb = load_workbook(output, data_only=True)
        po = wb['PO#_260823TP']
        self.assertEqual(po['R4'].value, 40)
        self.assertEqual(po['T4'].value, 4)
        self.assertEqual(po['H1'].value, 'Ttl Units Shipped 10/12-10/16 (via LAURA)')
        issues = json.loads(output.with_suffix('.issues.json').read_text())
        self.assertEqual(issues['区间']['开始'], '2026-10-17')
        self.assertEqual(issues['区间来源行数'], 1)

    def test_normalization_and_color_boundaries(self):
        self.assertEqual(report.po('PO#_CA260824TP'), 'CA260824')
        self.assertNotEqual(report.po('260824-1TP'), report.po('260824TP'))
        self.assertEqual(report.color_match('CloudFreeze Cooling Comfrter-ASH (92)', [(2, 'CLOUD'), (3, 'ASH')])[0], 3)
        self.assertEqual(report.color_match('Pillow OFF WHITE (20)', [(2, 'WHITE'), (3, 'OFF WHITE')])[0], 3)
        self.assertEqual(report.color_match('Pillow DUSTY ROSE (20)', [(2, 'ROSE'), (3, 'DUSTY ROSE')])[0], 3)
        self.assertIsNone(report.color_match('Pillow Red', [(2, 'Red'), (3, 'Red')])[0])


if __name__ == '__main__':
    unittest.main()
