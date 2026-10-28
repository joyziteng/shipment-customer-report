#!/usr/bin/env python3
"""Update PO detail worksheets without reserializing the rest of an XLSX file.

Inputs are read with openpyxl; writing is a selective, byte-preserving ZIP/XML
patch. Reconciliation is intentionally separate from the untouched TOTALS sheet.
"""
import argparse
import copy
import datetime as dt
import json
import math
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

FORWARDERS = ('LAURA', 'FLEXPORT')
MEASURES = ('units', 'percent', 'cbm')
SOURCE_TITLE = '单证部--工厂实际提货数据汇总'
SOURCE_HEADERS = {5: '款号', 6: '产品描述', 7: '外合同', 11: '数量',
                  14: '工厂提货日', 19: '总体积', 25: '货代反馈情况', 32: '货代'}


def po(value):
    """Only remove approved PO prefix, whitespace and final TP; keep punctuation."""
    text = re.sub(r'\s+', '', str(value or '').upper())
    text = re.sub(r'^PO#?_?', '', text)
    return re.sub(r'TP$', '', text)


def style(value):
    text = re.sub(r'[^A-Z0-9]', '', str(value or '').upper())
    return re.sub(r'TP$', '', text)


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def date_value(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%m/%d/%Y', '%Y.%m.%d'):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def interval_from_header(value, year, historical_year=None):
    text = str(value or '')
    if not re.search(r'\bSHIPPED', text, re.I):
        return None
    before = re.search(r'\bBEFORE\s+(\d{1,2})\s*/\s*(\d{1,2})(?:\s*/\s*(\d{4}))?', text, re.I)
    if before:
        if not before.group(3) and historical_year is None:
            raise ValueError(f'历史表头 {text!r} 的 BEFORE 缺少年份，请用 --historical-year 指定已确认年份')
        cutoff = dt.date(int(before.group(3) or historical_year), int(before.group(1)), int(before.group(2)))
        return dt.date.min, cutoff - dt.timedelta(days=1)
    iso = re.search(r'(\d{4}-\d{2}-\d{2})\s*(?:~|\bTO\b)\s*(\d{4}-\d{2}-\d{2})', text, re.I)
    if iso:
        a, b = (dt.date.fromisoformat(x) for x in iso.groups())
        if a > b:
            raise ValueError(f'日期表头起始晚于结束: {text!r}')
        return a, b
    match = re.search(r'(?<!\d)(\d{1,2})\s*/\s*(\d{1,2})\s*-\s*(\d{1,2})\s*/\s*(\d{1,2})(?!\d)', text)
    if match:
        am, ad, bm, bd = map(int, match.groups())
        a, b = dt.date(year - (am > bm), am, ad), dt.date(year, bm, bd)
        if a > b:
            raise ValueError(f'日期表头起始晚于结束: {text!r}')
        return a, b
    return None


def header_measure(value):
    text = str(value or '').strip().upper()
    if not re.search(r'\bSHIPPED', text):
        return None
    if re.match(r'^%', text):
        return 'percent'
    if re.match(r'^(?:(?:TTL|TOTAL)\s+)?UNITS\s+SHIPPED', text):
        return 'units'
    if re.match(r'^(?:(?:TTL|TOTAL)\s+)?CBM\s+SHIPPED', text):
        return 'cbm'
    return None


def header_forwarder(value):
    text = str(value or '').upper()
    found = [name for name in FORWARDERS if re.search(r'(?<![A-Z])' + name + r'(?![A-Z])', text)]
    return found[0] if len(found) == 1 else None


def color_match(description, candidates):
    """Match complete color tokens; prefer the last, then the most specific token."""
    text = str(description or '').upper().replace('CHESNUT', 'CHESTNUT')
    ranked = []
    for row, color in candidates:
        color = str(color or '').upper().replace('CHESNUT', 'CHESTNUT')
        parts = re.findall(r'[A-Z0-9]+', color)
        if not parts:
            continue
        # Optional separators allow BrownLeopard and Brown Leopard, without
        # matching CLOUD inside the product name CloudFreeze.
        pattern = r'(?<![A-Z0-9])' + r'[\s/_-]*'.join(map(re.escape, ''.join(parts))) + r'(?![A-Z0-9])'
        matches = list(re.finditer(pattern, text))
        if matches:
            matched = matches[-1]
            ranked.append((matched.end(), matched.start(), sum(map(len, parts)), row))
    if not ranked:
        return None, []
    # A complete multiword color ending later includes its shorter suffix, e.g.
    # OFF WHITE must win over WHITE even though both end at the same position.
    best = max((end, -start, size) for end, start, size, _ in ranked)
    winners = [row for end, start, size, row in ranked if (end, -start, size) == best]
    return (winners[0] if len(winners) == 1 else None), winners


def decimal(value):
    return Decimal(str(value))


def interval_label(start, end):
    """Match the customer workbook's M/D-M/D header convention; keep ISO across years."""
    if start.year == end.year:
        return f'{start.month}/{start.day}-{end.month}/{end.day}'
    return f'{start.isoformat()}~{end.isoformat()}'


def sheet_prefers_blank(item):
    """True when this sheet leaves zero shipments blank instead of writing 0."""
    ws, rows = item['ws'], item['rows']
    none_count = zero_count = 0
    for col in range(1, ws.max_column + 1):
        if header_measure(ws.cell(1, col).value) not in ('units', 'cbm'):
            continue
        for r in rows:
            value = item['cached'].cell(r, col).value
            if value is None:
                none_count += 1
            elif number(value) and value == 0:
                zero_count += 1
    return none_count > zero_count


def as_json_number(value):
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral() else float(value)
    return value


def sheet_parts(path):
    """Resolve sheet names through workbook relationships, not guessed filenames."""
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError('客户工作簿 ZIP 含重复条目，不能安全更新')
        if any(info.flag_bits & 1 for info in archive.infolist()):
            raise ValueError('客户工作簿 ZIP 含加密条目，不能安全更新')
        if any(name.startswith('_xmlsignatures/') for name in names):
            raise ValueError('客户工作簿含数字签名，修改会破坏签名，停止处理')
        wb = ET.fromstring(archive.read('xl/workbook.xml'))
        rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    by_id = {r.attrib['Id']: r for r in rels}
    result = {}
    for elem in wb.findall('.//{*}sheet'):
        rid = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        rel = by_id.get(rid)
        if rel is None or not rel.attrib.get('Type', '').endswith('/worksheet'):
            continue
        target = rel.attrib['Target']
        path_in_zip = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
        if path_in_zip not in names or not path_in_zip.startswith('xl/'):
            raise ValueError(f'无法安全定位工作表 XML: {elem.attrib.get("name")}')
        result[elem.attrib['name']] = path_in_zip
    return result


def read_customer(wb, cached, start, end, historical_year, issues, allow_overlap=False):
    """Identify all target sheets and resolve each sheet's exact or new columns."""
    sheets = []
    index = defaultdict(list)
    for ws in wb.worksheets:
        if ws.title == 'TOTALS':
            continue
        if not re.match(r'^\s*PO\s*#\s*[_\s]*(?:CA|UK)?\d{6}', ws.title, re.I):
            continue
        title_match = re.match(r'^\s*PO\s*#\s*[_\s]*((?:CA|UK)?\d{6}(?:-[A-Z0-9]+)?(?:TP)?)', ws.title, re.I)
        title_key = po(title_match.group(1))
        h = [str(ws.cell(1, c).value or '').strip().upper() for c in (1, 2, 3, 4, 5, 7)]
        if not (h[0].startswith('PO') and 'STYLE' in h[1] and 'DESCRIPTION' in h[2]
                and 'COLOR' in h[3] and 'UNIT' in h[4] and 'CBM' in h[5]):
            issues.append({'类型': '非标准PO明细页未处理', '页': ws.title})
            continue
        if ws.protection.sheet:
            raise ValueError(f'{ws.title} 受工作表保护，无法安全更新')
        rows = [r for r in range(2, ws.max_row + 1)
                if po(ws.cell(r, 1).value) and style(ws.cell(r, 2).value)]
        if not rows:
            issues.append({'类型': '明细页无可识别PO行', '页': ws.title})
            continue
        row_keys = {po(ws.cell(r, 1).value) for r in rows}
        if not (title_key in row_keys or any(k in row_keys for k in ('CA' + title_key, 'UK' + title_key))):
            issues.append({'类型': 'PO页标题与客户行不符，未处理', '页': ws.title,
                           '标题PO': title_key, '行PO': sorted(row_keys)})
            continue
        summary = max(rows) + 1
        if ws.cell(summary, 1).value or ws.cell(summary, 2).value:
            raise ValueError(f'{ws.title} 无法确认明细后的汇总行 {summary}')
        order_values = [cached[ws.title].cell(r, 5).value for r in rows]
        order_total = cached[ws.title].cell(summary, 5).value
        formula = str(ws.cell(summary, 5).value or '').upper()
        if not (('SUM(' in formula) or (number(order_total) and all(number(v) for v in order_values)
                                          and abs(sum(order_values) - order_total) < 1e-6)):
            raise ValueError(f'{ws.title}!E{summary} 无法确认汇总行，请检查布局')
        if any(ws.cell(r, 1).value and not style(ws.cell(r, 2).value) for r in range(2, summary)):
            issues.append({'类型': '客户行缺少款号', '页': ws.title,
                           '行': [r for r in range(2, summary) if ws.cell(r, 1).value and not style(ws.cell(r, 2).value)]})
        exact = defaultdict(list)
        templates = defaultdict(list)
        overlapped = []
        for cell in ws[1]:
            measure = header_measure(cell.value)
            if not measure:
                continue
            interval = interval_from_header(cell.value, end.year, historical_year)
            forwarder = header_forwarder(cell.value)
            if interval is None:
                raise ValueError(f'{ws.title}!{cell.coordinate} 出货表头日期无法识别: {cell.value!r}')
            if forwarder:
                templates[(forwarder, measure)].append(cell.column)
            if interval[0] <= end and start <= interval[1]:
                if interval != (start, end):
                    if not allow_overlap:
                        suggestion = (f'{interval[1] + dt.timedelta(days=1)}~{end}' if start <= interval[1] < end
                                      else f'{start}~{interval[0] - dt.timedelta(days=1)}' if start < interval[0] <= end
                                      else '拆分为不重叠的日期段')
                        raise ValueError(f'{ws.title}!{cell.coordinate} 历史区间 {interval[0]}~{interval[1]} 与请求 {start}~{end} 部分交叠，停止以免重复统计；可改用 {suggestion}，或确认后加 --allow-overlap 追加新列')
                    overlapped.append((cell.coordinate, interval))
                    continue
                if not forwarder:
                    raise ValueError(f'{ws.title}!{cell.coordinate} 对应区间的货代无法识别')
                exact[(forwarder, measure)].append(cell.column)
        if overlapped:
            if not allow_overlap:
                raise ValueError(f'{ws.title} 区间与已有列交叠')
            issues.append({'类型': '区间与已有列交叠，按新列追加', '页': ws.title,
                           '交叠单元格': [f'{coord} ({a}~{b})' for coord, (a, b) in overlapped],
                           '处理': '原列保留不动；本次新区间另追加列，交叠日期会在两组列中各出现一次'})
            exact.clear()
        if exact:
            missing = [f'{f}/{m}' for f in FORWARDERS for m in MEASURES if len(exact[(f, m)]) != 1]
            if missing:
                raise ValueError(f'{ws.title} 已有该区间，但货代数量/比例/体积列不完整或重复: {", ".join(missing)}')
            cols = {key: values[0] for key, values in exact.items()}
            mode = '原区间'
        else:
            first = ws.max_column + 1
            if first + 5 > 16384:
                raise ValueError(f'{ws.title} 追加日期列超过 Excel 列数上限')
            cols = {(f, m): first + i * 3 + j for i, f in enumerate(FORWARDERS) for j, m in enumerate(MEASURES)}
            mode = '新增区间'
        for key, col in cols.items():
            for merged in ws.merged_cells.ranges:
                if merged.min_col <= col <= merged.max_col and merged.min_row <= summary and merged.max_row >= 1:
                    raise ValueError(f'{ws.title} 目标列与合并区域 {merged} 交叠')
        item = {'ws': ws, 'cached': cached[ws.title], 'rows': rows, 'summary': summary,
                'columns': cols, 'templates': templates, 'mode': mode}
        sheets.append(item)
        for key in {po(ws.cell(r, 1).value) for r in rows}:
            index[key].append(item)
    for key, items in index.items():
        if len(items) > 1:
            issues.append({'类型': '重复PO明细页', 'PO': key, '页': [i['ws'].title for i in items],
                           '处理': '各页独立匹配更新；原始汇总按来源行计一次，TOTALS未修改'})
    return sheets, index


def read_source(path, start, end, index, issues):
    data = load_workbook(path, data_only=True, read_only=True)
    raw = load_workbook(path, data_only=False, read_only=True)
    try:
        if len(data.worksheets) < 2 or data.worksheets[1].title != SOURCE_TITLE:
            raise ValueError(f'出货表第二页必须是 {SOURCE_TITLE}')
        ws, ws_raw = data.worksheets[1], raw.worksheets[1]
        for col, name in SOURCE_HEADERS.items():
            if name not in str(ws.cell(1, col).value or ''):
                raise ValueError(f'出货表 {get_column_letter(col)} 列表头不符，预期包含 {name}')
        records = []
        undated = defaultdict(list)
        aggregates = defaultdict(lambda: {'行号': [], '数量': Decimal(0), '体积': Decimal(0), '数量不可用行': [], '体积不可用行': []})
        for r, (row, raw_row) in enumerate(zip(ws.iter_rows(min_row=2, max_col=32),
                                               ws_raw.iter_rows(min_row=2, max_col=32)), start=2):
            key = po(row[6].value)
            day = date_value(row[13].value)
            if day is None:
                if key in index:
                    undated[key].append(r)
                continue
            if not start <= day <= end:
                continue
            st, desc = style(row[4].value), str(row[5].value or '')
            if not key:
                issues.append({'类型': '区间内外合同缺失', '出货表行': r, '款号': st})
                continue
            f_raw = str(row[31].value or '').strip()
            forwarder = f_raw.upper()
            q, cbm = row[10].value, row[18].value
            ok_q, ok_cbm = number(q) and q >= 0, number(cbm) and cbm >= 0
            record = {'row': r, 'po': key, 'style': st, 'description': desc, 'date': str(day),
                      'forwarder': forwarder, 'q': decimal(q) if ok_q else None,
                      'cbm': decimal(cbm) if ok_cbm else None, 'note': str(row[24].value or '')}
            records.append(record)
            group = aggregates[(key, f_raw or '（空）')]
            group['行号'].append(r)
            if ok_q:
                group['数量'] += decimal(q)
            else:
                group['数量不可用行'].append(r)
            if ok_cbm:
                group['体积'] += decimal(cbm)
            else:
                group['体积不可用行'].append(r)
            if record['note'] and key in index:
                issues.append({'类型': '货代反馈需核对', '出货表行': r, 'PO': key, '备注': record['note'],
                               '数量': q, '体积': cbm, '处理': '按N/AF/K/S缓存计入，不根据备注猜测修正'})
            if forwarder not in FORWARDERS and key in index:
                issues.append({'类型': '货代无法识别', '出货表行': r, 'PO': key, '货代': f_raw,
                               '数量': q, '体积': cbm, '处理': '不猜测货代；原始汇总保留'})
            if not st and key in index:
                issues.append({'类型': '款号缺失', '出货表行': r, 'PO': key})
            if (not ok_q or not ok_cbm) and key in index:
                issues.append({'类型': '数量或体积缓存不可用', '出货表行': r, 'PO': key, '数量': str(q), '体积': str(cbm),
                               '源公式': [str(raw_row[c - 1].value) for c in (11, 19)]})
            elif ok_q and q > 0 and cbm == 0 and key in index:
                dims = [row[c - 1].value for c in (16, 17, 18)]
                issues.append({'类型': '体积为零需核对', '出货表行': r, 'PO': key, '数量': q, '体积': cbm,
                               '箱规P到R': dims, '源体积公式': str(raw_row[18].value),
                               '处理': '按缓存0计入；检查箱规缺失及体积'})
            if key not in index:
                issues.append({'类型': '来源PO无客户明细页', '出货表行': r, 'PO': key,
                               '处理': '不创建工作表；原始汇总保留'})
        for key, row_numbers in sorted(undated.items()):
            issues.append({'类型': '提货日期不可识别（未计入任何区间）', 'PO': key,
                           '出货表行': row_numbers,
                           '处理': '这些行没有可识别的工厂提货日，需人工确认归属日期'})
        return records, aggregates
    finally:
        data.close()
        raw.close()


def build_updates(item, records, start, end, issues):
    ws, cached, rows, summary, cols = (item[k] for k in ('ws', 'cached', 'rows', 'summary', 'columns'))
    keys = {po(ws.cell(r, 1).value) for r in rows}
    relevant = [rec for rec in records if rec['po'] in keys]
    lookup = defaultdict(list)
    for r in rows:
        lookup[(po(ws.cell(r, 1).value), style(ws.cell(r, 2).value))].append((r, ws.cell(r, 4).value))
    placed = defaultdict(lambda: {'units': Decimal(0), 'cbm': Decimal(0), 'records': []})
    blocked = set()  # (customer row, forwarder, measure)
    for rec in relevant:
        candidates = lookup[(rec['po'], rec['style'])]
        target, winners = color_match(rec['description'], candidates)
        if target is None:
            uncertain_rows = [r for r, _color in candidates]
            if not rec['style']:
                uncertain_rows = [r for r in rows if po(ws.cell(r, 1).value) == rec['po']]
            for r in uncertain_rows:
                for f in FORWARDERS:
                    for m in ('units', 'cbm'):
                        blocked.add((r, f, m))
            issues.append({'类型': '款号颜色无法唯一匹配', '页': ws.title, '出货表行': rec['row'],
                           'PO': rec['po'], '款号': rec['style'], '描述': rec['description'],
                           '候选行': winners, '同款行': [r for r, _ in candidates],
                           '处理': '同PO同款原区间值保留；来源汇总独立列示'})
            continue
        if rec['forwarder'] not in FORWARDERS:
            for f in FORWARDERS:
                for m in ('units', 'cbm'):
                    blocked.add((target, f, m))
            issues.append({'类型': '明细货代无法归集', '页': ws.title, '客户行': target, '出货表行': rec['row'],
                           'PO': rec['po'], '处理': '保留该行各货代已有区间值，来源汇总列示'})
            continue
        f = rec['forwarder']
        group = placed[(target, f)]
        group['records'].append(rec)
        for m, field in (('units', 'q'), ('cbm', 'cbm')):
            if rec[field] is None:
                blocked.add((target, f, m))
            else:
                group[m] += rec[field]
    changes, differences, columns_json = {}, [], {}
    label = interval_label(start, end)
    blank_zero = sheet_prefers_blank(item)
    for f in FORWARDERS:
        columns_json[f] = {m: get_column_letter(cols[(f, m)]) for m in MEASURES}
        for m in MEASURES:
            col = cols[(f, m)]
            if item['mode'] == '新增区间':
                names = {'units': 'Ttl Units', 'percent': '% To Ttl Units', 'cbm': 'CBM'}
                changes[(1, col)] = ('text', f'{names[m]} Shipped {label} (via {f})', template_style(item, f, m, 1))
        for r in rows:
            group = placed[(r, f)]
            q = group['units'] if (r, f, 'units') not in blocked else None
            cbm = group['cbm'] if (r, f, 'cbm') not in blocked else None
            order = cached.cell(r, 5).value
            percentage = q / decimal(order) if q is not None and number(order) and order > 0 else None
            if q is not None and percentage is None:
                issues.append({'类型': '订单量不可用，比例保留', '页': ws.title, '行': r, '订单量': str(order), '货代': f})
            for m, value in {'units': q, 'percent': percentage, 'cbm': cbm}.items():
                if value is None:
                    continue
                col = cols[(f, m)]
                old = cached.cell(r, col).value if item['mode'] == '原区间' else None
                if m == 'percent' and item['mode'] == '原区间' and ws.cell(r, col).data_type == 'f':
                    continue  # ratio formulas stay live and recalculate from the updated units
                if value == 0 and (old is None if item['mode'] == '原区间' else blank_zero):
                    continue  # keep the sheet's blank-for-zero convention
                changes[(r, col)] = ('number', value, template_style(item, f, m, r))
                if item['mode'] == '原区间' and (not number(old) or abs(decimal(old) - value) > Decimal('0.000000001')):
                    differences.append({'单元格': f'{get_column_letter(col)}{r}', 'PO': po(ws.cell(r, 1).value),
                                        '货代': f, '指标': m, '原缓存': old, '新值': as_json_number(value),
                                        '原公式或值': str(ws.cell(r, col).value),
                                        '来源行': [rec['row'] for rec in group['records']]})
        # Sum only the values actually visible, including any preserved cells.
        for m in ('units', 'cbm'):
            col = cols[(f, m)]
            total, missing = Decimal(0), []
            for r in rows:
                entry = changes.get((r, col))
                if entry is not None:
                    total += decimal(entry[1])
                    continue
                if item['mode'] == '原区间':
                    old = cached.cell(r, col).value
                    if number(old):
                        total += decimal(old)
                    elif old is not None:
                        missing.append(r)
                # 新增区间中未写入的零值按空白处理，不计入缺失
            if missing:
                issues.append({'类型': '区间汇总不完整', '页': ws.title, '货代': f, '指标': m, '行': missing,
                               '处理': '汇总仅包括可确认数值，须结合原始汇总核对'})
            changes[(summary, col)] = ('number', total, template_style(item, f, m, summary))
            old_total = cached.cell(summary, col).value if item['mode'] == '原区间' else None
            if item['mode'] == '原区间' and (not number(old_total) or abs(decimal(old_total) - total) > Decimal('0.000000001')):
                differences.append({'单元格': f'{get_column_letter(col)}{summary}', 'PO': '本页合计',
                                    '货代': f, '指标': m, '原缓存': old_total, '新值': as_json_number(total),
                                    '原公式或值': str(ws.cell(summary, col).value), '来源行': []})
            if m == 'units':
                total_units = total
        all_orders = [cached.cell(r, 5).value for r in rows]
        if all(number(v) and v >= 0 for v in all_orders) and sum(all_orders) > 0:
            total_order = sum((decimal(v) for v in all_orders), Decimal(0))
            percent_col = cols[(f, 'percent')]
            ratio = total_units / total_order
            if not (item['mode'] == '原区间' and ws.cell(summary, percent_col).data_type == 'f'):
                changes[(summary, percent_col)] = ('number', ratio, template_style(item, f, 'percent', summary))
            old_ratio = cached.cell(summary, percent_col).value if item['mode'] == '原区间' else None
            if item['mode'] == '原区间' and (summary, percent_col) in changes and (
                    not number(old_ratio) or abs(decimal(old_ratio) - ratio) > Decimal('0.000000001')):
                differences.append({'单元格': f'{get_column_letter(percent_col)}{summary}', 'PO': '本页合计',
                                    '货代': f, '指标': 'percent', '原缓存': old_ratio, '新值': as_json_number(ratio),
                                    '原公式或值': str(ws.cell(summary, percent_col).value), '来源行': []})
        else:
            issues.append({'类型': '汇总订单量不可用，比例保留', '页': ws.title, '货代': f})
    if blocked:
        issues.append({'类型': '部分明细保留原值', '页': ws.title,
                       '单元格': [f'{get_column_letter(cols[(f, m)])}{r}' for r, f, m in sorted(blocked)],
                       '处理': '无法唯一归集或缓存缺失的指标保留原值；另一有效指标仍可更新'})
    allocated = defaultdict(lambda: {'units': Decimal(0), 'cbm': Decimal(0), 'units_rows': [], 'cbm_rows': []})
    for (r, f), group in placed.items():
        for rec in group['records']:
            for m, field in (('units', 'q'), ('cbm', 'cbm')):
                if (r, f, m) not in blocked and rec[field] is not None:
                    dest = allocated[(rec['po'], f)]
                    dest[m] += rec[field]
                    dest[m + '_rows'].append(rec['row'])
    groups = defaultdict(list)
    for rec in relevant:
        groups[(rec['po'], rec['forwarder'] or '（空）')].append(rec)
    reconciliation = []
    complete_rows = set()
    for (key, f), group in sorted(groups.items()):
        dest = allocated[(key, f)]
        total_q = sum((rec['q'] for rec in group if rec['q'] is not None), Decimal(0))
        total_cbm = sum((rec['cbm'] for rec in group if rec['cbm'] is not None), Decimal(0))
        complete_rows.update(set(dest['units_rows']) & set(dest['cbm_rows']))
        reconciliation.append({'PO': key, '货代': f, '来源行': [rec['row'] for rec in group],
                               '来源数量': as_json_number(total_q), '来源体积': as_json_number(total_cbm),
                               '已归集数量': as_json_number(dest['units']), '未归集数量': as_json_number(total_q - dest['units']),
                               '已归集体积': as_json_number(dest['cbm']), '未归集体积': as_json_number(total_cbm - dest['cbm']),
                               '数量已归集来源行': dest['units_rows'], '体积已归集来源行': dest['cbm_rows'],
                               '数量不可用行': [rec['row'] for rec in group if rec['q'] is None],
                               '体积不可用行': [rec['row'] for rec in group if rec['cbm'] is None]})
    return changes, {'页': ws.title, '状态': ws.sheet_state, '处理': item['mode'], '列': columns_json,
                     '来源行数': len(relevant), '来源行': [rec['row'] for rec in relevant],
                     '已完整归集来源行': sorted(complete_rows),
                     '未完整归集来源行': sorted(rec['row'] for rec in relevant if rec['row'] not in complete_rows),
                     '保留原值单元格': [f'{get_column_letter(cols[(f, m)])}{r}' for r, f, m in sorted(blocked)],
                     '按PO货代核对': reconciliation, '覆盖差异': differences}


def template_style(item, forwarder, measure, row):
    ws = item['ws']
    target_col = item['columns'][(forwarder, measure)]
    if item['mode'] == '原区间':
        return None  # existing XML already carries the original style ID
    candidates = item['templates'].get((forwarder, measure)) or item['templates'].get((FORWARDERS[0], measure)) or []
    exemplar = ws.cell(row, candidates[-1] if candidates else 5)
    if exemplar._style is None:
        return 0
    # style_id may create a new normalized ID not present in the preserved styles.xml.
    for index, original in enumerate(ws.parent._cell_styles):
        if tuple(original) == tuple(exemplar._style):
            return index
    raise ValueError(f'{ws.title}!{exemplar.coordinate} 样式不在原工作簿样式表中')


def sheet_xml_patch(xml, changes):
    """Replace only selected cell elements, preserving untouched XML byte-for-byte."""
    prefix_match = re.search(rb'<((?:[\w.-]+:)?)(?:worksheet)\b', xml)
    if not prefix_match:
        raise ValueError('无法识别工作表 XML 根元素')
    prefix = prefix_match.group(1)
    qrow, qc = re.escape(prefix + b'row'), re.escape(prefix + b'c')
    row_pattern = re.compile(rb'<(' + qrow + rb')\b([^>]*?)(?:/>|>(.*?)</\1\s*>)', re.S)
    cell_pattern = re.compile(rb'<(' + qc + rb')\b([^>]*?)(?:/>|>(.*?)</\1\s*>)', re.S)
    by_row = defaultdict(dict)
    for (r, c), value in changes.items():
        by_row[r][c] = value
    found_rows = set()

    def replace_row(match):
        attrs, body = match.group(2), match.group(3)
        rid = re.search(rb'\br\s*=\s*["\'](\d+)["\']', attrs)
        if not rid or int(rid.group(1)) not in by_row:
            return match.group(0)
        r = int(rid.group(1))
        found_rows.add(r)
        body = body or b''
        pending = dict(by_row[r])
        pieces, pos = [], 0
        for cell in cell_pattern.finditer(body):
            pieces.append(body[pos:cell.start()])
            coord = re.search(rb'\br\s*=\s*["\']([A-Za-z]+)(\d+)["\']', cell.group(2))
            if not coord:
                raise ValueError(f'工作表第 {r} 行包含无坐标单元格')
            col = column_index_from_string(coord.group(1).decode().upper())
            for add in sorted(c for c in pending if c < col):
                pieces.append(make_cell(prefix, r, add, pending.pop(add), None, None))
            if col in pending:
                pieces.append(make_cell(prefix, r, col, pending.pop(col), cell.group(2), cell.group(3)))
            else:
                pieces.append(cell.group(0))
            pos = cell.end()
        pieces.append(body[pos:])
        for col in sorted(pending):
            pieces.append(make_cell(prefix, r, col, pending[col], None, None))
        max_col = max(by_row[r])
        spans = re.search(rb'\bspans\s*=\s*(["\'])(\d+):(\d+)\1', attrs)
        if spans and int(spans.group(3)) < max_col:
            replacement = b'spans=' + spans.group(1) + spans.group(2) + b':' + str(max_col).encode() + spans.group(1)
            attrs = attrs[:spans.start()] + replacement + attrs[spans.end():]
        return b'<' + prefix + b'row' + attrs + b'>' + b''.join(pieces) + b'</' + prefix + b'row>'

    patched = row_pattern.sub(replace_row, xml)
    if found_rows != set(by_row):
        raise ValueError(f'目标行在原始 XML 中不存在: {sorted(set(by_row) - found_rows)}')
    max_col = max(c for _, c in changes)
    max_row = max(r for r, _ in changes)
    dimension = re.compile(rb'(<'+re.escape(prefix)+rb'dimension\b[^>]*\bref\s*=\s*["\'])([^"\']+)(["\'])')
    def expand_dimension(match):
        raw = match.group(2).decode()
        found = re.findall(r'\$?([A-Za-z]+)\$?(\d+)', raw)
        if not found:
            raise ValueError(f'工作表 dimension 无法识别: {raw}')
        old_col = max(column_index_from_string(col.upper()) for col, _ in found)
        old_row = max(int(r) for _, r in found)
        return match.group(1) + f'A1:{get_column_letter(max(max_col, old_col))}{max(max_row, old_row)}'.encode() + match.group(3)
    if not dimension.search(patched):
        raise ValueError('工作表 XML 缺少 dimension，无法安全扩展范围')
    return dimension.sub(expand_dimension, patched, count=1) if max_col > 0 and any(r == 1 for r, _ in changes) else patched


def make_cell(prefix, row, col, change, attrs, body):
    kind, value, style_id = change
    if attrs is None:
        attrs = f' r="{get_column_letter(col)}{row}"'.encode()
        if style_id:
            attrs += f' s="{style_id}"'.encode()
    else:
        # Shared/array formulas span multiple cells and cannot safely be split.
        if body and re.search(rb'<(?:[\w.-]+:)?f\b[^>]*\bt\s*=\s*["\'](?:shared|array|dataTable)["\']', body):
            raise ValueError(f'{get_column_letter(col)}{row} 属于共享或数组公式，不能局部覆盖')
        if body and re.search(rb'<(?:[\w.-]+:)?(?:extLst|ext)\b', body):
            raise ValueError(f'{get_column_letter(col)}{row} 含单元格扩展元数据，不能安全覆盖')
        attrs = re.sub(rb'\s+t\s*=\s*(["\']).*?\1', b'', attrs)
    if kind == 'text':
        attrs += b' t="inlineStr"'
        contents = b'<' + prefix + b'is><' + prefix + b't>' + escape(str(value)).encode('utf-8') + b'</' + prefix + b't></' + prefix + b'is>'
    else:
        if not value.is_finite():
            raise ValueError(f'{get_column_letter(col)}{row} 数值无效')
        contents = b'<' + prefix + b'v>' + format(value, 'f').encode() + b'</' + prefix + b'v>'
    return b'<' + prefix + b'c' + attrs + b'>' + contents + b'</' + prefix + b'c>'


def patch_archive(customer, destination, patches):
    """Copy untouched local ZIP entries verbatim and regenerate changed sheets.

    The central directory and offsets are rebuilt, but non-target local headers,
    compressed bytes, descriptors, ZIP metadata and uncompressed content remain
    byte-identical. No orphan sheet records are left behind.
    """
    with zipfile.ZipFile(customer) as original, open(customer, 'rb') as raw, zipfile.ZipFile(destination, 'w') as archive:
        archive.comment = original.comment
        infos = original.infolist()
        ordered = sorted(infos, key=lambda info: info.header_offset)
        ends = {info.header_offset: ordered[i + 1].header_offset if i + 1 < len(ordered) else original.start_dir
                for i, info in enumerate(ordered)}
        if ordered and ordered[0].header_offset != 0:
            raise ValueError('客户 ZIP 含前置数据，无法安全复制')
        for info in ordered:
            if info.filename in patches:
                archive.writestr(copy.copy(info), patches[info.filename])
                continue
            new_info = copy.copy(info)
            new_info.header_offset = archive.fp.tell()
            raw.seek(info.header_offset)
            segment = raw.read(ends[info.header_offset] - info.header_offset)
            if len(segment) < 30 or segment[:4] != b'PK\x03\x04':
                raise ValueError(f'ZIP 本地条目无效: {info.filename}')
            archive.fp.write(segment)
            archive.start_dir = archive.fp.tell()
            archive.filelist.append(new_info)
            archive.NameToInfo[new_info.filename] = new_info
            archive._didModify = True
    with zipfile.ZipFile(customer) as original, zipfile.ZipFile(destination) as result:
        if set(original.namelist()) != set(result.namelist()) or result.testzip() is not None:
            raise ValueError('输出 ZIP 结构校验失败')
        for name in original.namelist():
            if name not in patches and original.read(name) != result.read(name):
                raise ValueError(f'非目标 ZIP 条目发生变化: {name}')
            if name in patches and result.read(name) != patches[name]:
                raise ValueError(f'目标工作表 XML 写入校验失败: {name}')


def totals_constant_cells(totals, start, end, historical_year, issues):
    """Warn when TOTALS cells for this interval are constants, so they cannot auto-sum."""
    columns = []
    for cell in totals[1]:
        if not header_measure(cell.value) or not header_forwarder(cell.value):
            continue
        try:
            interval = interval_from_header(cell.value, end.year, historical_year)
        except ValueError:
            continue
        if interval == (start, end):
            columns.append((header_forwarder(cell.value), header_measure(cell.value), cell.column))
    if not columns:
        return
    constants = []
    for r in range(2, totals.max_row + 1):
        key = po(totals.cell(r, 1).value)
        if not key:
            continue
        for forwarder, measure, col in columns:
            value = totals.cell(r, col).value
            if value is not None and not (isinstance(value, str) and value.lstrip().startswith('=')):
                constants.append({'行': r, 'PO': key, '货代': forwarder, '指标': measure,
                                  '单元格': f'{get_column_letter(col)}{r}', '当前值': value})
    if constants:
        issues.append({'类型': 'TOTALS对应日期列是固定值，不会自动汇总', '数量': len(constants),
                       '明细': constants, '处理': 'TOTALS按用户要求不修改；如需汇总，请在 Excel 中人工处理'})


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description='仅更新客户PO明细页，生成独立核对报告；TOTALS原样保留')
    parser.add_argument('--customer', required=True)
    parser.add_argument('--shipment', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--days', type=int, help='以昨天为结束日，包含昨天的自然日数')
    parser.add_argument('--start', help='起始日期 YYYY-MM-DD（含）')
    parser.add_argument('--end', help='结束日期 YYYY-MM-DD（含）')
    parser.add_argument('--today', help='配合 --days 复核时指定今天 YYYY-MM-DD')
    parser.add_argument('--historical-year', type=int, help='已确认的无年份 BEFORE 表头年份')
    parser.add_argument('--allow-overlap', action='store_true',
                        help='允许与已有日期列部分重叠：原列不动，本次区间追加新列')
    args = parser.parse_args(argv)
    if args.days is not None:
        if args.start or args.end or args.days < 1:
            parser.error('--days 必须大于0，且不能与 --start/--end 同时使用')
        end = (dt.date.fromisoformat(args.today) if args.today else dt.date.today()) - dt.timedelta(days=1)
        start = end - dt.timedelta(days=args.days - 1)
    else:
        if not args.start or not args.end or args.today:
            parser.error('请指定 --start 和 --end，或指定 --days；--today 仅配合 --days')
        start, end = dt.date.fromisoformat(args.start), dt.date.fromisoformat(args.end)
        if start > end:
            parser.error('--start 不能晚于 --end')
    return args, start, end


def main(argv=None):
    args, start, end = parse_args(argv)
    customer, shipment, output = (Path(p).expanduser().resolve() for p in (args.customer, args.shipment, args.output))
    report_path = output.with_suffix('.issues.json')
    if output in (customer, shipment) or report_path in (customer, shipment) or output.exists() or report_path.exists():
        raise ValueError('输出及核对报告不能覆盖输入或已有文件')
    if output.suffix.lower() != '.xlsx':
        raise ValueError('输出必须是 .xlsx')
    paths = sheet_parts(customer)
    wb = load_workbook(customer, data_only=False)
    cached = load_workbook(customer, data_only=True)
    try:
        if 'TOTALS' not in wb or 'TOTALS' not in paths:
            raise ValueError('找不到 TOTALS 工作表')
        issues = []
        sheets, index = read_customer(wb, cached, start, end, args.historical_year, issues, args.allow_overlap)
        totals_constant_cells(wb['TOTALS'], start, end, args.historical_year, issues)
        records, aggregates = read_source(shipment, start, end, index, issues)
        details, patches = [], {}
        with zipfile.ZipFile(customer) as archive:
            for item in sheets:
                changes, detail = build_updates(item, records, start, end, issues)
                name = item['ws'].title
                if name not in paths or paths[name] == paths['TOTALS']:
                    raise ValueError(f'无法安全定位 PO 工作表 {name}')
                patches[paths[name]] = sheet_xml_patch(archive.read(paths[name]), changes)
                details.append(detail)
        agg_json = [{'PO': key, '货代原文': f, '来源行数': len(group['行号']),
                     '来源行': group['行号'], '数量': as_json_number(group['数量']),
                     '体积': as_json_number(group['体积']), '数量不可用行': group['数量不可用行'],
                     '体积不可用行': group['体积不可用行']}
                    for (key, f), group in sorted(aggregates.items())]
        zeroed = []
        for detail in details:
            hits = [d for d in detail['覆盖差异']
                    if number(d.get('原缓存')) and d['原缓存'] != 0 and d['新值'] in (0, 0.0)]
            if hits:
                zeroed.append({'页': detail['页'], '原值被归零单元格数': len(hits), '示例': hits[:5]})
        report = {'区间': {'开始': str(start), '结束': str(end), '天数': (end - start).days + 1},
                  '区间来源行数': len(records), '来源PO货代汇总': agg_json, '明细页核对': details,
                  '原值非零但按源表归零': zeroed,
                  '待确认数量': len(issues), '待确认': issues,
                  '说明': ['TOTALS及所有非目标ZIP条目保持原样，TOTALS不会自动反映本次更新或新增日期。',
                           '原始PO货代汇总按来源行统计一次；重复明细页分别核对，不将页面总额相加。',
                           '历史待出货及其他公式缓存保持原样，可能不反映本次变化；打开或重算需另行核查。',
                           '原值非零但按源表归零表示该指标在指定期间、该货代下没有出货；如与业务记忆不符，请核对来源行。',
                           '未归集、备注、体积异常和覆盖差异需人工核对；未确认前结果为草稿。']}
        report_text = json.dumps(report, ensure_ascii=False, indent=2, default=str)
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(prefix='.shipment-', suffix='.xlsx', dir=output.parent, delete=False) as temp:
            temp_path = Path(temp.name)
        try:
            patch_archive(customer, temp_path, patches)
            # Reopen to catch malformed XML and numerical encodings before publishing.
            verified = load_workbook(temp_path, data_only=True, read_only=False)
            try:
                for detail in details:
                    _ = verified[detail['页']]['A1'].value
            finally:
                verified.close()
            os.link(temp_path, output)  # atomic creation, never overwrite
            try:
                with report_path.open('x', encoding='utf-8') as report_file:
                    report_file.write(report_text)
            except Exception:
                output.unlink()
                raise
        finally:
            temp_path.unlink(missing_ok=True)
        print(json.dumps({'output': str(output), 'issues': str(report_path), 'start': str(start), 'end': str(end),
                          'source_rows': len(records), 'issue_count': len(issues), 'sheets': len(details)}, ensure_ascii=False))
    finally:
        wb.close()
        cached.close()


if __name__ == '__main__':
    try:
        main()
    except (ValueError, OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        print(f'错误: {exc}', file=sys.stderr)
        sys.exit(2)
